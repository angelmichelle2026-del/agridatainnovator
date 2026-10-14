#!/usr/bin/env python3
"""Import all tables from the supplied NISR SAS 2026 Season A workbook.

Requires openpyxl. The importer retains each worksheet's tabular observations,
creates normalized Table 9 district land use, and derives area-weighted
province/city crop yields from Tables 12, 13 and 15.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "artifacts/agri-lens-ai/public/data/nisr-sas-2026-season-a-tables.xlsx"
TABLES_OUT = ROOT / "artifacts/agri-lens-ai/src/data/nisr-sas-2026-season-a-tables.json"
CROP_VIS_OUT = ROOT / "artifacts/agri-lens-ai/src/data/crop-visuals-2026a.json"
DISTRICT_OUT = ROOT / "artifacts/agri-lens-ai/src/data/district-land-2026a.json"
DISTRICT_CSV = ROOT / "artifacts/agri-lens-ai/public/data/district-land-2026-season-a.csv"
ADMIN_AREAS_IN = ROOT / "artifacts/agri-lens-ai/src/data/rwanda-admin-areas.json"
PROVINCE_YIELD_OUT = ROOT / "artifacts/agri-lens-ai/src/data/province-yields-2026a.json"
SOURCE_ID = "nisr-sas-2026-season-a-workbook"
SOURCE_URL = "https://www.statistics.gov.rw/statistical-publications/agriculture-environment/seasonal-agricultural-survey-season-2026"

# Header row numbers and first data row for each worksheet. Multi-row headings
# are combined while preserving blank cells as nulls and numeric source values.
LAYOUT = {
    1: ((3,), 4), 2: ((3,), 4), 3: ((3,), 4), 4: ((4,), 5),
    5: ((3, 4), 5), 6: ((3, 4), 5), 7: ((3, 4), 5), 8: ((3, 4), 5),
    9: ((3,), 4), 10: ((4, 5), 6), 11: ((3,), 4), 12: ((3,), 4),
    13: ((3,), 4), 14: ((3,), 4), 15: ((4,), 5), 16: ((3,), 4),
    17: ((4, 5), 6), 18: ((4,), 5), 19: ((4,), 5), 20: ((4, 5), 6),
    21: ((3,), 4), 22: ((3, 4), 6), 23: ((3,), 4), 24: ((3, 4), 5),
    25: ((3, 4), 5), 26: ((3,), 4), 27: ((3,), 4), 28: ((3,), 4),
    29: ((3, 4), 5), 30: ((3,), 4), 31: ((3, 4), 5), 32: ((3, 4), 5),
    33: ((3,), 4), 34: ((4, 5), 6), 35: ((3,), 4),
}

DISTRICT_FIELDS = [
    "total_land_000ha", "agricultural_land_000ha", "agricultural_land_pct",
    "arable_land_000ha", "physical_cultivated_land_000ha",
    "seasonal_crop_area_000ha", "permanent_crop_area_000ha",
    "temporary_fallow_area_000ha", "temporary_meadow_pasture_000ha",
    "permanent_pasture_000ha",
]


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def source_row(ws, row_number: int) -> bool:
    return any(
        isinstance(cell.value, str) and cell.value.strip().lower().startswith("source:")
        for cell in ws[row_number]
    )


def title_for(ws, number: int) -> str:
    pattern = re.compile(rf"^\s*Table\s*{number}\s*:", re.IGNORECASE)
    for row in ws.iter_rows():
        values = [str(cell.value).strip() for cell in row if cell.value is not None]
        if any(pattern.match(value) for value in values):
            return " ".join(values)
    return f"Table {number}"


def merged_header_value(ws, row_number: int, column_number: int) -> Any:
    cell = ws.cell(row_number, column_number)
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_row <= row_number <= merged.max_row
            and merged.min_col <= column_number <= merged.max_col
        ):
            return ws.cell(merged.min_row, merged.min_col).value
    return cell.value


def unique_headers(ws, header_rows: tuple[int, ...], columns: list[int]) -> list[str]:
    headers = []
    for column in columns:
        parts = []
        for row_number in header_rows:
            value = merged_header_value(ws, row_number, column)
            text = str(value).strip() if value is not None else ""
            if text and text not in parts:
                parts.append(text)
        header = " · ".join(parts) or f"Column {column}"
        headers.append(header)
    seen: dict[str, int] = {}
    for index, header in enumerate(headers):
        seen[header] = seen.get(header, 0) + 1
        if seen[header] > 1:
            headers[index] = f"{header} ({seen[header]})"
    return headers


def workbook_tables(wb) -> list[dict[str, Any]]:
    result = []
    for number in range(1, 36):
        ws = wb[f"Table {number}"]
        header_rows, data_start = LAYOUT[number]
        data_end = next(
            (row for row in range(data_start, ws.max_row + 1) if source_row(ws, row)),
            ws.max_row + 1,
        )
        row_numbers = [r for r in range(data_start, data_end) if any(c.value is not None for c in ws[r])]
        relevant_rows = list(header_rows) + row_numbers
        columns = [
            column
            for column in range(1, ws.max_column + 1)
            if any(ws.cell(row, column).value is not None for row in relevant_rows)
        ]
        headers = unique_headers(ws, header_rows, columns)
        if number == 35 and headers and headers[0].startswith("Column "):
            headers[0] = "District"
        rows = [
            [json_value(ws.cell(row, column).value) for column in columns]
            for row in row_numbers
        ]
        result.append({
            "id": f"table-{number:02d}",
            "sheet_name": ws.title,
            "number": number,
            "title": title_for(ws, number),
            "header_rows": list(header_rows),
            "row_count": len(rows),
            "headers": headers,
            "rows": rows,
            "source_id": SOURCE_ID,
        })
    return result


def district_land_table9(ws) -> dict[str, Any]:
    # Table 9 uses (,000 Ha). Its first column is blank due to workbook layout;
    # District is column B, followed by the ten land-use measures in C:L.
    observations = []
    for row_number in range(4, 35):
        values = [ws.cell(row_number, col).value for col in range(2, 13)]
        if not values or values[0] is None:
            continue
        record = {
            "district": str(values[0]).strip(),
            "year": 2026,
            "season": "A",
            "source_id": SOURCE_ID,
        }
        record.update({field: json_value(value) for field, value in zip(DISTRICT_FIELDS, values[1:])})
        observations.append(record)
    if len(observations) != 31:
        raise ValueError(f"Expected 31 Table 9 rows (30 districts + National); got {len(observations)}")
    return {
        "source": {
            "id": SOURCE_ID,
            "publisher": "National Institute of Statistics of Rwanda (NISR)",
            "title": "SAS 2026 Season A statistical tables (workbook)",
            "period": "September 2025 to February 2026",
            "tables": ["Table 9: Agricultural land use per district"],
            "source_url": SOURCE_URL,
            "workbook_file": "nisr-sas-2026-season-a-tables.xlsx",
            "unit_note": "Land-use areas are thousand hectares (000 ha); percentages and workbook precision are retained.",
            "geography_note": "30 districts plus the workbook's National row.",
            "limitations": "This extract is Table 9 only; other workbook tables are available in the full table browser. The workbook was supplied for this task and its worksheets cite NISR SAS 2026.",
        },
        "fields": DISTRICT_FIELDS,
        "observations": observations,
    }


def normalized_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def numeric_value(value: Any) -> bool:
    return isinstance(value, (int, float)) and math.isfinite(value)


def province_yield_data(
    tables: list[dict[str, Any]], administrative_areas: dict[str, Any]
) -> dict[str, Any]:
    by_number = {table["number"]: table for table in tables}
    harvested_area = by_number[12]
    district_yield = by_number[13]
    production = by_number[15]
    excluded_groups = {
        "cereals",
        "tubersandroots",
        "banana",
        "legumesandpulses",
        "bean",
        "vegetablesandfruits",
    }
    crops = [
        crop
        for crop in district_yield["headers"][1:]
        if normalized_header(crop) not in excluded_groups
    ]

    def header_index(table: dict[str, Any], crop: str) -> int:
        target = normalized_header(crop)
        for index, header in enumerate(table["headers"]):
            if normalized_header(header) == target:
                return index
        raise ValueError(f"Table {table['number']} has no column for {crop!r}")

    def rows_by_name(table: dict[str, Any]) -> dict[str, list[Any]]:
        return {str(row[0]).strip(): row for row in table["rows"]}

    area_rows = rows_by_name(harvested_area)
    yield_rows = rows_by_name(district_yield)
    production_rows = rows_by_name(production)
    entities = administrative_areas["entities"]
    district_names = [
        district for entity in entities for district in entity["districts"]
    ]
    if len(district_names) != 30 or len(set(district_names)) != 30:
        raise ValueError("Administrative mapping must contain 30 unique districts")

    observations = []
    national_benchmark = []
    for crop in crops:
        area_index = header_index(harvested_area, crop)
        yield_index = header_index(district_yield, crop)
        production_index = header_index(production, crop)
        grouped = []
        national_area = 0.0
        national_production = 0.0

        for entity in entities:
            area_total = 0.0
            production_total = 0.0
            for district in entity["districts"]:
                area = area_rows[district][area_index]
                crop_production = production_rows[district][production_index]
                reported_yield = yield_rows[district][yield_index]
                if not all(
                    numeric_value(value)
                    for value in (area, crop_production, reported_yield)
                ):
                    raise ValueError(
                        f"Missing {crop} area, production or yield for {district}"
                    )
                if area < 0 or crop_production < 0:
                    raise ValueError(f"Negative source value for {crop} in {district}")
                if area == 0:
                    if crop_production != 0:
                        raise ValueError(
                            f"Positive {crop} production with zero harvested area in {district}"
                        )
                else:
                    calculated_yield = crop_production * 1000 / area
                    if not math.isclose(
                        calculated_yield,
                        reported_yield,
                        rel_tol=1e-9,
                        abs_tol=1e-6,
                    ):
                        raise ValueError(
                            f"Table 13 yield mismatch for {crop} in {district}: "
                            f"{calculated_yield} vs {reported_yield} kg/ha"
                        )
                area_total += area
                production_total += crop_production
            national_area += area_total
            national_production += production_total
            grouped.append(
                {
                    "province": entity["name"],
                    "crop": crop,
                    "harvested_area_ha": area_total,
                    "production_mt": production_total,
                    "yield_kg_ha": (
                        production_total * 1000 / area_total
                        if area_total > 0
                        else None
                    ),
                }
            )

        national_row = yield_rows.get("National")
        if not national_row or not numeric_value(national_row[yield_index]):
            raise ValueError(f"Table 13 has no numeric national yield for {crop}")
        derived_national_yield = (
            national_production * 1000 / national_area if national_area > 0 else None
        )
        if derived_national_yield is None or not math.isclose(
            derived_national_yield,
            national_row[yield_index],
            rel_tol=1e-9,
            abs_tol=1e-6,
        ):
            raise ValueError(f"District-weighted {crop} yield does not match Table 13 national value")
        observations.extend(grouped)
        national_benchmark.append(
            {
                "crop": crop,
                "yield_kg_ha": national_row[yield_index],
                "harvested_area_ha": national_area,
                "production_mt": national_production,
            }
        )

    return {
        "source": {
            "id": SOURCE_ID,
            "source_url": SOURCE_URL,
            "period": "2026 Season A",
            "table_numbers": [12, 13, 15],
            "method": (
                "Province/city yield is derived per crop as the sum of district "
                "production (MT) divided by summed harvested crop area (Ha), "
                "multiplied by 1,000 to report kg/ha. This area-weights the district yields."
            ),
            "interpretation_note": (
                "Province/city values are derived aggregations, not figures "
                "separately published by NISR. The national comparison is the "
                "Table 13 national yield. This is a single-season snapshot, not a multi-season trend."
            ),
        },
        "crops": crops,
        "provinces": [
            {"name": entity["name"], "rw_name": entity["rw_name"]}
            for entity in entities
        ],
        "national_benchmark": national_benchmark,
        "observations": observations,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    if len(wb.worksheets) != 35:
        raise ValueError(f"Expected 35 worksheets; found {len(wb.worksheets)}")

    tables = workbook_tables(wb)
    administrative_areas = json.loads(ADMIN_AREAS_IN.read_text(encoding="utf-8"))
    province_yields = province_yield_data(tables, administrative_areas)
    tables_data = {
        "source": {
            "id": SOURCE_ID,
            "publisher": "National Institute of Statistics of Rwanda (NISR)",
            "title": "SAS 2026 Season A statistical tables",
            "period": "September 2025 to February 2026",
            "source_url": SOURCE_URL,
            "workbook_file": "nisr-sas-2026-season-a-tables.xlsx",
            "worksheet_count": len(tables),
            "provenance_note": "Workbook supplied for this task; its worksheets attribute the tables to NISR SAS 2026. The linked NISR page is the Season A release landing page.",
        },
        "tables": tables,
    }
    crop_visual_data = {
        "source": {
            "id": SOURCE_ID,
            "source_url": SOURCE_URL,
            "period": "2026 Season A",
            "table_numbers": [6, 15],
            "description": "Compact chart data: same-season national crop production in Tables 6 and 2026A district production in Table 15.",
        },
        "national_production": tables[5],
        "district_production": tables[14],
    }
    district_data = district_land_table9(wb["Table 9"])
    TABLES_OUT.parent.mkdir(parents=True, exist_ok=True)
    CROP_VIS_OUT.parent.mkdir(parents=True, exist_ok=True)
    DISTRICT_OUT.parent.mkdir(parents=True, exist_ok=True)
    PROVINCE_YIELD_OUT.parent.mkdir(parents=True, exist_ok=True)
    DISTRICT_CSV.parent.mkdir(parents=True, exist_ok=True)
    TABLES_OUT.write_text(json.dumps(tables_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    CROP_VIS_OUT.write_text(json.dumps(crop_visual_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    PROVINCE_YIELD_OUT.write_text(json.dumps(province_yields, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    DISTRICT_OUT.write_text(json.dumps(district_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with DISTRICT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["district", "year", "season", *DISTRICT_FIELDS, "source_id", "source_url"])
        writer.writeheader()
        for row in district_data["observations"]:
            writer.writerow({**row, "source_url": SOURCE_URL})
    print(f"Imported {len(tables)} tables, {len(district_data['observations'])} 2026 Season A district land-use rows, and {len(province_yields['observations'])} area-weighted province/crop yields.")
    print(f"Tables: {TABLES_OUT}\nCrop charts: {CROP_VIS_OUT}\nProvince yield charts: {PROVINCE_YIELD_OUT}\nDistrict JSON: {DISTRICT_OUT}\nDistrict CSV: {DISTRICT_CSV}")


if __name__ == "__main__":
    main()
