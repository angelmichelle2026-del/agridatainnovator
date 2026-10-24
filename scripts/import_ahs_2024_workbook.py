#!/usr/bin/env python3
"""Import the supplied NISR Agricultural Household Survey (AHS) 2024 workbook.

The normalized browser catalogue keeps each table's worksheet row numbers,
Excel column letters, displayed values (including cached formula results), and
original workbook unchanged for full-fidelity download.
"""
from __future__ import annotations

import argparse
import json
import math
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "artifacts/agri-lens-ai/public/data/AHS_2024Tables.xlsx"
OUTPUT = ROOT / "artifacts/agri-lens-ai/src/data/ahs-2024-tables.json"
SOURCE_ID = "nisr-ahs-2024-workbook"
SOURCE_URL = "/data/AHS_2024Tables.xlsx"
EXPECTED_TABLES = 56


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value
    return str(value)


def table_number(sheet_name: str) -> int:
    match = re.fullmatch(r"\s*Table\s*(\d+)\s*", sheet_name, re.IGNORECASE)
    if not match:
        raise ValueError(f"Unexpected AHS table worksheet name: {sheet_name!r}")
    return int(match.group(1))


def table_title(ws, number: int) -> str:
    pattern = re.compile(rf"^\s*Table\s*{number}\s*:", re.IGNORECASE)
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and pattern.match(cell.value):
                return " ".join(cell.value.split())
    raise ValueError(f"Table {number} has no matching title row")


def load_tables(workbook_path: Path) -> dict[str, Any]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"AHS 2024 workbook not found: {workbook_path}")

    formulas_wb = load_workbook(workbook_path, data_only=False, read_only=False)
    values_wb = load_workbook(workbook_path, data_only=True, read_only=False)
    numbers = sorted(
        table_number(name)
        for name in values_wb.sheetnames
        if re.fullmatch(r"\s*Table\s*\d+\s*", name, re.IGNORECASE)
    )
    if numbers != list(range(1, EXPECTED_TABLES + 1)):
        raise ValueError(
            f"Expected AHS tables 1-{EXPECTED_TABLES}; found {numbers!r}"
        )
    if "Contents" not in values_wb.sheetnames:
        raise ValueError("AHS 2024 workbook has no Contents worksheet")

    tables = []
    uncached_formulas = []
    for number in numbers:
        sheet_name = next(
            name for name in values_wb.sheetnames
            if re.fullmatch(rf"\s*Table\s*{number}\s*", name, re.IGNORECASE)
        )
        ws = values_wb[sheet_name]
        formula_ws = formulas_wb[sheet_name]
        max_row = max(
            (cell.row for row in ws.iter_rows() for cell in row if cell.value is not None),
            default=0,
        )
        max_col = max(
            (cell.column for row in ws.iter_rows() for cell in row if cell.value is not None),
            default=0,
        )
        min_col = min(
            (cell.column for row in ws.iter_rows() for cell in row if cell.value is not None),
            default=1,
        )
        if not max_row or not max_col:
            raise ValueError(f"Table {number} worksheet is empty")
        columns = [get_column_letter(index) for index in range(min_col, max_col + 1)]
        rows = []
        for row_number in range(1, max_row + 1):
            values = [ws.cell(row_number, column).value for column in range(min_col, max_col + 1)]
            if not any(value is not None for value in values):
                continue
            cells = []
            for column, value in zip(range(min_col, max_col + 1), values):
                original = formula_ws.cell(row_number, column)
                if original.data_type == "f" and value is None:
                    uncached_formulas.append(f"{sheet_name}!{original.coordinate}")
                cells.append(json_value(value))
            rows.append({"excel_row": row_number, "cells": cells})

        tables.append({
            "id": f"ahs-2024-table-{number:02d}",
            "number": number,
            "sheet_name": sheet_name,
            "title": table_title(ws, number),
            "columns": columns,
            "row_count": len(rows),
            "rows": rows,
            "source_id": SOURCE_ID,
        })

    if uncached_formulas:
        raise ValueError(
            "Formula cells are missing cached results: " + ", ".join(uncached_formulas[:20])
        )

    contents = values_wb["Contents"]
    contents_text = " ".join(
        str(cell.value).strip()
        for row in contents.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.strip()
    )
    if "AHS 2024" not in contents_text:
        raise ValueError("Contents sheet does not identify the AHS 2024 survey")

    source = {
        "id": SOURCE_ID,
        "publisher": "National Institute of Statistics of Rwanda (NISR)",
        "title": "Agricultural Household Survey (AHS) 2024 — Main Tables",
        "period": "Agricultural Household Survey 2024; some livestock tables cover the 2023/24 agricultural year",
        "source_url": SOURCE_URL,
        "workbook_file": "AHS_2024Tables.xlsx",
        "worksheet_count": len(tables),
        "citation": "National Institute of Statistics of Rwanda (NISR), Agricultural Household Survey 2024 report, August 2025.",
        "provenance_note": "Values are presented from the supplied NISR AHS 2024 workbook. Worksheet cells, source rows, table-specific periods, units and cached formula results are retained; no new estimates are derived.",
    }
    return {"source": source, "tables": tables}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    data = load_tables(args.workbook.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print(
        f"Imported {len(data['tables'])} AHS 2024 tables "
        f"({sum(table['row_count'] for table in data['tables'])} non-empty source rows) "
        f"to {args.output}"
    )


if __name__ == "__main__":
    main()
